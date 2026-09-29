"""Structured OOXML renderer for a format-neutral report document."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .model import ReportDocument


def _hyperlink(paragraph, text: str, url: str) -> None:
    relation = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "520978")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _set_cell(cell, value: object, *, bold: bool = False) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(value))
    run.bold = bold
    run.font.name = "DejaVu Sans"
    run.font.size = Pt(9)


def render_docx(document: ReportDocument) -> bytes:
    """Generate a proper WordprocessingML document fully in memory."""
    word = Document()
    section = word.sections[0]
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    styles = word.styles
    styles["Normal"].font.name = "DejaVu Sans"
    styles["Normal"].font.size = Pt(10)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        styles[style_name].font.name = "DejaVu Sans"
        styles[style_name].font.color.rgb = RGBColor(0x52, 0x09, 0x78)

    heading = word.add_heading("SourceHealth", 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = word.add_paragraph("Repository Health Report")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.size = Pt(18)
    repo = word.add_paragraph(document.repository)
    repo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    score_table = word.add_table(rows=2, cols=2)
    score_table.style = "Light Shading Accent 1"
    _set_cell(score_table.cell(0, 0), "Repo Health Score", bold=True)
    _set_cell(score_table.cell(0, 1), "Coverage", bold=True)
    _set_cell(score_table.cell(1, 0), document.health_score if document.health_score is not None else "NO_DATA", bold=True)
    _set_cell(score_table.cell(1, 1), f"{document.coverage_percent}%" if document.coverage_percent is not None else "NO_DATA", bold=True)

    word.add_heading("Analysis metadata", level=1)
    metadata = [
        ("Repository", document.repository),
        ("Analysis ID", document.analysis_id),
        ("Analysis date/time", document.completed_at),
        ("Analysis status", document.status),
        ("Scoring policy", document.scoring_policy_version),
    ]
    table = word.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"
    for label, value in metadata:
        if value:
            cells = table.add_row().cells
            _set_cell(cells[0], label, bold=True)
            _set_cell(cells[1], value)
    cells = table.add_row().cells
    _set_cell(cells[0], "Canonical SourceCraft URL", bold=True)
    cells[1].text = ""
    _hyperlink(cells[1].paragraphs[0], document.canonical_url, document.canonical_url)

    word.add_heading("Category summary", level=1)
    table = word.add_table(rows=1, cols=5)
    table.style = "Light Shading Accent 1"
    for cell, label in zip(table.rows[0].cells, ("Category", "Weight", "Score", "Availability", "Explanation"), strict=True):
        _set_cell(cell, label, bold=True)
    for category in document.categories:
        cells = table.add_row().cells
        values = (category.label, f"{category.weight}%", category.score if category.score is not None else "NO_DATA",
                  category.availability, category.explanation or "—")
        for cell, value in zip(cells, values, strict=True):
            _set_cell(cell, value)

    word.add_heading("Category details", level=1)
    for category in document.categories:
        word.add_heading(f"{category.label} — {category.weight}%", level=2)
        word.add_paragraph(
            f"Score: {category.score if category.score is not None else 'NO_DATA'} · "
            f"Availability: {category.availability}"
        )
        word.add_paragraph(category.explanation or "Explanation is not available in the stored analysis.")
        if category.evidence_refs:
            word.add_paragraph("Evidence: " + ", ".join(category.evidence_refs))

    word.add_heading("Recommendations", level=1)
    if document.recommendations:
        for item in document.recommendations:
            word.add_heading(f"P{item.priority} — {item.title}", level=2)
            word.add_paragraph(item.description)
            word.add_paragraph("Action: " + (item.action or "—"))
            if item.evidence_refs:
                word.add_paragraph("Evidence: " + ", ".join(item.evidence_refs))
    else:
        word.add_paragraph("No recommendations are present in the stored analysis.")

    word.add_heading("Checks, evidence and limitations", level=1)
    for check in document.checks:
        word.add_heading(f"{check.name} — {check.availability}", level=2)
        word.add_paragraph(f"Source: {check.source or 'not specified'} · Findings: {check.finding_count}")
        for evidence in check.evidence:
            paragraph = word.add_paragraph(style="List Bullet")
            paragraph.add_run(f"{evidence.identifier}: {evidence.summary} ({evidence.source})")
            if evidence.url:
                paragraph.add_run(" — ")
                _hyperlink(paragraph, evidence.url, evidence.url)
            elif evidence.reference:
                paragraph.add_run(" — " + evidence.reference)
    word.add_paragraph(
        "NO_DATA means that the stored analysis does not contain enough confirmed data; it is not a zero score. "
        "The document displays the stored analysis and does not recalculate Health Score."
    )
    footer = section.footer.paragraphs[0]
    footer.text = "SourceHealth · Repository Health Report"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    output = BytesIO()
    word.save(output)
    return output.getvalue()
