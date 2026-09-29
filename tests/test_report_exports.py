"""PDF, DOCX and Markdown exports share one stored report truth."""

import io
import unittest
import zipfile
from contextlib import nullcontext
from types import SimpleNamespace
from unittest.mock import Mock, patch
from uuid import UUID

from docx import Document
from pypdf import PdfReader

from sourcehealth.api.routers.analyses import docx, markdown, pdf
from sourcehealth.application.services import ServiceError
from sourcehealth.markdown import render_markdown
from sourcehealth.reports.docx import render_docx
from sourcehealth.reports.model import ReportDocument
from sourcehealth.reports.pdf import render_pdf

ANALYSIS_ID = UUID("12345678-1234-5678-1234-567812345678")


def report_fixture():
    categories = {
        "documentation": {"score": 88, "availability": "available", "explanation": "Документация проверена.",
                          "evidence_refs": ["docs-readme"]},
        "cicd": {"score": 75, "availability": "available", "explanation": "CI настроен.", "evidence_refs": []},
        "security": {"score": None, "availability": "no_data", "explanation": "Нет данных SourceCraft AppSec.",
                     "evidence_refs": []},
        "activity": {"score": 70, "availability": "available", "explanation": "Есть свежие изменения.",
                     "evidence_refs": []},
        "issues": {"score": 60, "availability": "partial", "explanation": "Данные доступны частично.",
                   "evidence_refs": []},
        "code_health": {"score": 92, "availability": "available", "explanation": "Локальные проверки завершены.",
                        "evidence_refs": []},
    }
    return {
        "schema_version": "3.0",
        "repository": {
            "organization_slug": "команда",
            "repository_slug": "проект",
            "canonical_url": "https://sourcecraft.dev/team/project",
        },
        "completed_at": "2026-09-29T09:00:00+00:00",
        "scoring_policy_version": "mvp-score-v1.2",
        "health_score": 79.5,
        "category_scores": categories,
        "recommendations": [{
            "id": "rec-docs",
            "priority": 1,
            "title": "Улучшить документацию",
            "description": "Добавить описание эксплуатации.",
            "suggested_action": "Дополнить README и runbook.",
            "evidence_refs": ["docs-readme"],
        }],
        "checks": {
            "documentation": {
                "availability": "available",
                "source": "sourcehealth_local",
                "findings": [],
                "evidence": [{
                    "id": "docs-readme",
                    "summary": "README обнаружен",
                    "source": "repository_metadata",
                    "reference": "README.md",
                    "url": "https://sourcecraft.dev/team/project/blob/main/README.md",
                }],
            },
        },
        "internal_secret": "PAT-MUST-NEVER-LEAK",
    }


def document():
    return ReportDocument.from_report(report_fixture(), analysis_id=ANALYSIS_ID, status="partial")


def docx_text(payload: bytes) -> str:
    word = Document(io.BytesIO(payload))
    values = [paragraph.text for paragraph in word.paragraphs]
    for table in word.tables:
        values.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(values)


class ReportRendererTests(unittest.TestCase):
    def test_all_formats_share_canonical_values_and_exclude_unknown_secrets(self):
        model = document()
        md = render_markdown(report_fixture())
        pdf_bytes = render_pdf(model)
        docx_bytes = render_docx(model)
        pdf_text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf_bytes)).pages)
        word_text = docx_text(docx_bytes)

        self.assertTrue(pdf_bytes.startswith(b"%PDF"))
        self.assertGreater(len(pdf_bytes), 10_000)
        self.assertTrue(docx_bytes.startswith(b"PK"))
        with zipfile.ZipFile(io.BytesIO(docx_bytes)) as archive:
            self.assertIn("word/document.xml", archive.namelist())
        for content in (md, pdf_text, word_text):
            normalized = content.replace("\\", "")
            self.assertIn("команда/проект", normalized)
            self.assertIn("79.5", normalized)
            self.assertIn("mvp-score-v1.2", normalized)
            self.assertIn("documentation", normalized.lower())
            self.assertIn("NO_DATA", normalized)
            self.assertIn("Улучшить документацию", normalized)
            self.assertNotIn("PAT-MUST-NEVER-LEAK", normalized)
        self.assertIn("README обнаружен", word_text)
        self.assertGreaterEqual(len(PdfReader(io.BytesIO(pdf_bytes)).pages), 2)

    def test_invalid_report_fails_without_rendering_internal_payload(self):
        with self.assertRaisesRegex(ValueError, "schema 3.0"):
            ReportDocument.from_report({"schema_version": "internal", "token": "secret"})


class ReportRouteTests(unittest.TestCase):
    def request(self, run):
        sessions = Mock(return_value=nullcontext(Mock()))
        return SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(sessions=sessions))), run

    def test_completed_export_responses_have_binary_contracts_and_filenames(self):
        request, run = self.request(SimpleNamespace(
            id=ANALYSIS_ID, status="completed", results=report_fixture(),
        ))
        with patch("sourcehealth.api.routers.analyses.public_run", return_value=run):
            md_response = markdown(ANALYSIS_ID, request)
            pdf_response = pdf(ANALYSIS_ID, request)
            docx_response = docx(ANALYSIS_ID, request)

        self.assertEqual(md_response.status_code, 200)
        self.assertEqual(md_response.media_type, "text/markdown; charset=utf-8")
        self.assertIn(".md", md_response.headers["content-disposition"])
        self.assertEqual(pdf_response.media_type, "application/pdf")
        self.assertIn(".pdf", pdf_response.headers["content-disposition"])
        self.assertTrue(pdf_response.body.startswith(b"%PDF"))
        self.assertEqual(
            docx_response.media_type,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn(".docx", docx_response.headers["content-disposition"])
        self.assertTrue(docx_response.body.startswith(b"PK"))

    def test_all_exports_preserve_report_not_ready(self):
        for status in ("queued", "collecting", "analyzing", "scoring"):
            request, run = self.request(SimpleNamespace(id=ANALYSIS_ID, status=status, results=report_fixture()))
            with patch("sourcehealth.api.routers.analyses.public_run", return_value=run):
                for endpoint in (markdown, pdf, docx):
                    with self.subTest(status=status, endpoint=endpoint.__name__), self.assertRaises(ServiceError) as raised:
                        endpoint(ANALYSIS_ID, request)
                    self.assertEqual(raised.exception.code, "report_not_ready")
                    self.assertEqual(raised.exception.status, 409)


if __name__ == "__main__":
    unittest.main()
